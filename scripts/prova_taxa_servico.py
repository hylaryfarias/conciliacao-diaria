"""Prova de cada par cupom × venda da máquina apontado como serviço não baixado.

    python3 scripts/prova_taxa_servico.py saida/isaias_taxa_servico_2026-10-01_a_18.xlsx

Para cada caso verifica, com as vendas de todos os caixas da loja e todas as modalidades da máquina:
- se o terminal da venda é usado só por aquele caixa (pelos casamentos exatos do período);
- quantos minutos separam a venda na máquina do fechamento do cupom;
- se existe outro pagamento no PDV, livre (não casado com venda própria), de valor igual ao da máquina, no turno;
- se existe outra venda na máquina, livre, de valor igual ao do cupom, no turno.
Acrescenta as colunas na própria planilha.
"""
import glob
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from concilia_pix import carrega_lojas  # noqa: E402
from pdv_pix import HORA_VIRADA, ler_pdv  # noqa: E402
from portal_sicredi import ler_portais  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent


def main():
    arq = sys.argv[1]
    lojas = carrega_lojas()
    b = ler_pdv(str(RAIZ / "entrada" / "pdv_cupons_*.xlsx"))
    b["dt"] = pd.to_datetime(b["Data"].astype(str) + " " + b["Hora"]); b.loc[b["dt"].dt.hour < HORA_VIRADA, "dt"] += pd.Timedelta(days=1)
    b["Cupom"] = b["Cupom"].astype(str); b["Caixa"] = b["Caixa"].astype(str).str.strip()
    portais = ler_portais(glob.glob(str(RAIZ / "entrada" / "portal" / "*.xlsx")), lojas)
    px = pd.concat([pd.read_excel(f, sheet_name="PDV x maquininha") for f in glob.glob(str(RAIZ / "saida" / "concilia_pix_2*.xlsx"))])
    cc = pd.concat([pd.read_excel(f, sheet_name="Maquininha x PDV") for mod in ("credito", "debito") for f in glob.glob(str(RAIZ / "saida" / f"concilia_{mod}_2*.xlsx"))])
    xl = pd.ExcelFile(arq)
    saida = {}
    for aba in xl.sheet_names:
        I = xl.parse(aba)
        filial = I["Loja"].iloc[0]
        bb = b[b["Filial"] == filial]
        m = next(df for (f, ec), df in portais.items() if f == filial)
        m = m[m["Status"].str.upper().isin({"APROVADA", "AUTORIZADA"})]
        p1 = px[(px["Filial"] == filial) & (px["Status"] == "CONCILIADO")]; c1 = cc[(cc["Filial"] == filial) & (cc["Status"] == "CONCILIADO")]
        casados = {(str(x["Caixa"]).strip(), str(x["Cupom"]), int(round(float(x["Vl. pagamento"]) * 100))) for _, x in p1.iterrows()}
        casados |= {(str(x["Caixa"]).strip(), str(x["Cupom"]).replace(".0", ""), int(round(float(x["Valor Bruto"]) * 100))) for _, x in c1.iterrows()}
        aut_casadas = set(p1["Autorização"].astype(str)) | set(c1["Autorização"].astype(str))
        tc = pd.concat([p1.groupby(["Terminal", "Caixa"]).size(), c1.groupby(["Terminal", "Caixa"]).size()]).groupby(level=[0, 1]).sum().unstack(fill_value=0)
        cols = []
        for _, r in I.iterrows():
            d0 = pd.to_datetime(r["Data"], dayfirst=True).date()
            dtm = pd.Timestamp(f"{d0} {r['Hora máquina']}"); dtm += pd.Timedelta(days=1) if dtm.hour < HORA_VIRADA else pd.Timedelta(0)
            dtp = pd.Timestamp(f"{d0} {r['Hora PDV']}"); dtp += pd.Timedelta(days=1) if dtp.hour < HORA_VIRADA else pd.Timedelta(0)
            vm = int(round(float(r["Valor máquina"]) * 100)); vp = int(round(float(r["Valor PDV"]) * 100))
            t = str(r["Terminal"]); usos = tc.loc[t] if t in tc.index else None
            if usos is not None and usos.sum() > 0:
                cx_term = f"{usos.idxmax()[-6:]} em {usos.max()} de {usos.sum()} casamentos exatos"
                exclusivo = "sim" if usos.max() == usos.sum() and usos.idxmax().strip() == str(r["Caixa"]).strip() else "não"
            else:
                cx_term, exclusivo = "sem casamento exato no período", "?"
            j = bb[(bb["dt"] >= dtm - pd.Timedelta(minutes=30)) & (bb["dt"] <= dtm + pd.Timedelta(hours=3)) & (bb["cent"] == vm)]
            livres = [x for _, x in j.iterrows() if (x["Caixa"], x["Cupom"], x["cent"]) not in casados]
            k = m[(m["dt"] >= dtp - pd.Timedelta(hours=3)) & (m["dt"] <= dtp + pd.Timedelta(minutes=30)) & (m["cent"] == vp)]
            livres_m = [x for _, x in k.iterrows() if str(x["Cód. de autorização"]) not in aut_casadas]
            cols.append({"Terminal usado por": cx_term, "Terminal exclusivo deste caixa": exclusivo,
                         "Minutos máquina → fechamento": round((dtp - dtm).total_seconds() / 60),
                         "Outro pagamento PDV livre = valor máquina (turno)": "; ".join(f'{x["Caixa"][-6:]} {x["Desc. pagam."]} {x["dt"]:%d/%m %H:%M}' for x in livres) or "nenhum",
                         "Outra venda máq livre = valor cupom (turno)": "; ".join(f'{x["dt"]:%d/%m %H:%M} {x["Modalidade"]} {x["Número do terminal"]}' for x in livres_m) or "nenhuma",
                         "Prova": "fechada" if exclusivo == "sim" and not livres and not livres_m else "ver observação"})
        saida[aba] = pd.concat([I, pd.DataFrame(cols)], axis=1)
    with pd.ExcelWriter(arq, engine="openpyxl") as xw:
        for aba, df in saida.items():
            df.to_excel(xw, sheet_name=aba, index=False); ws = xw.book[aba]
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), 1):
                ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, max(len(str(c.value or "")) for c in col) + 2))
            ws.freeze_panes = "A2"
            print(aba, df["Prova"].value_counts().to_dict())
            print(df[df["Prova"] != "fechada"][["Cupom", "Data", "Valor PDV", "Valor máquina", "Terminal", "Terminal exclusivo deste caixa", "Outro pagamento PDV livre = valor máquina (turno)", "Outra venda máq livre = valor cupom (turno)"]].to_string(index=False))


if __name__ == "__main__":
    main()
