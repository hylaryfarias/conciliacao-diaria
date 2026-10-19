"""Serviço não baixado: cupons em que a máquina cobrou o PDV + 10% e o cupom fechou sem o item da taxa.

    python3 scripts/taxa_servico.py --de 2026-10-01 --ate 2026-10-18 [--analitico entrada/pdv_analitico_*.xlsx]

Lê as abas "Serviço não baixado" de saida/concilia_pix_*, concilia_credito_*, concilia_debito_* (todas as lojas),
e grava saida/taxa_servico_<de>_a_<ate>.xlsx com: Lançamentos +10%, Por loja, Por dia, e para o ISAIAS 01 as
colunas de operador, atendente e mesa vindas do relatório analítico (quando houver).
"""
import glob
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from python_calamine import CalamineWorkbook

RAIZ = Path(__file__).resolve().parent.parent


def le(padrao):
    fs = glob.glob(str(RAIZ / "saida" / padrao))
    return pd.concat([pd.read_excel(f, sheet_name="Serviço não baixado") for f in fs], ignore_index=True) if fs else pd.DataFrame()


def analitico(arqs):
    """Uma linha por cupom: operador, atendente, mesa e se o cupom tem o item TAXA DE SERVICO."""
    partes = []
    for a in arqs:
        wb = CalamineWorkbook.from_path(a)
        rows = wb.get_sheet_by_name(wb.sheet_names[0]).to_python()
        h = next(i for i, r in enumerate(rows) if "Operador" in r and "Produto" in r)
        cols = [str(c) for c in rows[h]]
        # 'Total' aparece duas vezes: renomeia a segunda
        seen = {}
        for i, c in enumerate(cols):
            if c in seen:
                cols[i] = f"{c} item"
            seen[c] = 1
        partes.append(pd.DataFrame(rows[h + 1:], columns=cols))
    if not partes:
        return pd.DataFrame()
    d = pd.concat(partes, ignore_index=True)
    d = d[d["Cupom"].astype(str).str.strip() != ""]
    d["Cupom"] = d["Cupom"].astype(str).str.replace(".0", "", regex=False)
    d["taxa"] = d["Produto"].astype(str).str.upper().str.contains("TAXA DE SERV")
    d["Preço"] = pd.to_numeric(d["Preço"], errors="coerce").fillna(0)
    g = d.groupby(["Filial", "Cupom"]).agg(**{
        "Operador do caixa": ("Operador", "first"),
        "Atendente": ("Nome do atendente", "first"),
        "Mesa": ("Nr. mesa/comanda", "first"),
        "Item taxa no cupom": ("taxa", "any"),
    }).reset_index()
    tx = d[d["taxa"]].groupby(["Filial", "Cupom"])["Preço"].sum().rename("Taxa no cupom (R$)").reset_index()
    g = g.merge(tx, on=["Filial", "Cupom"], how="left")
    g["Item taxa no cupom"] = g["Item taxa no cupom"].map({True: "sim", False: "não"})
    return g


def main():
    args = sys.argv[1:]
    de = pd.Timestamp(args[args.index("--de") + 1]).date() if "--de" in args else None
    ate = pd.Timestamp(args[args.index("--ate") + 1]).date() if "--ate" in args else None
    ana = []
    if "--analitico" in args:
        i = args.index("--analitico")
        ana = [a for parte in args[i + 1].split(",") for a in (glob.glob(parte) or [parte])]
    else:
        ana = glob.glob(str(RAIZ / "entrada" / "pdv_analitico_*.xlsx"))

    linhas = []
    px = le("concilia_pix_2*.xlsx")
    for _, x in px.iterrows():
        pdv = float(x["Vl. pagamento"]); sv = float(x.get("Serviço (R$)", 0) or 0)
        linhas.append({"Forma": "PIX", "Filial": x["Filial"], "Data": pd.to_datetime(x["Data"]).date(), "Cupom": str(x["Cupom"]).replace(".0", ""), "Caixa": x["Caixa"],
                       "Hora PDV": str(x["Hora"])[:5], "Valor PDV": round(pdv, 2), "Hora máquina": str(x.get("Hora maq", ""))[:5], "Valor máquina": round(pdv + sv, 2),
                       "Terminal": x.get("Terminal", ""), "Autorização": x.get("Autorização", ""), "Serviço (R$)": round(sv, 2)})
    for mod, nome in (("credito", "Crédito"), ("debito", "Débito")):
        cx = le(f"concilia_{mod}_2*.xlsx")
        for _, x in cx.iterrows():
            bruto = float(x["Valor Bruto"]); sv = float(x.get("Serviço (R$)", 0) or 0)
            linhas.append({"Forma": nome, "Filial": x["Filial"], "Data": pd.to_datetime(x["Dia PDV"] if pd.notna(x.get("Dia PDV")) else x["Data"]).date(), "Cupom": str(x["Cupom"]).replace(".0", ""), "Caixa": x["Caixa"],
                           "Hora PDV": str(x.get("Hora PDV", ""))[:5], "Valor PDV": round(bruto - sv, 2), "Hora máquina": str(x["Hora maq"])[:5], "Valor máquina": round(bruto, 2),
                           "Terminal": x.get("Terminal", ""), "Autorização": x.get("Autorização", ""), "Serviço (R$)": round(sv, 2)})
    L = pd.DataFrame(linhas)
    if de:
        L = L[L["Data"] >= de]
    if ate:
        L = L[L["Data"] <= ate]
    L["Dia"] = pd.to_datetime(L["Data"]).dt.day_name().map({"Monday": "seg", "Tuesday": "ter", "Wednesday": "qua", "Thursday": "qui", "Friday": "sex", "Saturday": "sáb", "Sunday": "dom"})
    L = L.sort_values(["Filial", "Data", "Hora PDV"]).reset_index(drop=True)

    por_loja = L.groupby("Filial").agg(Cupons=("Cupom", "count"), **{"Serviço (R$)": ("Serviço (R$)", "sum")}).reset_index()
    for f in ("PIX", "Crédito", "Débito"):
        por_loja[f] = por_loja["Filial"].map(L[L["Forma"] == f].groupby("Filial")["Cupom"].count()).fillna(0).astype(int)
    tot_pdv = {}
    for f in glob.glob(str(RAIZ / "saida" / "concilia_*_2*.xlsx")):
        if "banco" in f:
            continue
        r = pd.read_excel(f, sheet_name="Resumo loja x dia")
        col = [c for c in r.columns if c.endswith("PDV (R$)")][0]
        for k, v in r.groupby("Filial")[col].sum().items():
            tot_pdv[k] = tot_pdv.get(k, 0) + v
    por_loja["Vendas no período (R$)"] = por_loja["Filial"].map(tot_pdv).round(2)
    por_loja["Serviço / vendas"] = (por_loja["Serviço (R$)"] / por_loja["Vendas no período (R$)"]).round(4)
    por_loja = por_loja.sort_values("Serviço (R$)", ascending=False)
    por_dia = L.groupby(["Filial", "Data", "Dia"]).agg(Cupons=("Cupom", "count"), **{"Serviço (R$)": ("Serviço (R$)", "sum")}).reset_index()

    A = analitico(ana)
    I = L[L["Filial"].str.startswith("ISAIAS 01")].copy()
    if len(A):
        I = I.merge(A, on=["Filial", "Cupom"], how="left")
    I_op = I.groupby("Operador do caixa").agg(Cupons=("Cupom", "count"), **{"Sobra (R$)": ("Serviço (R$)", "sum")}).reset_index().sort_values("Sobra (R$)", ascending=False) if "Operador do caixa" in I.columns else pd.DataFrame()

    out = RAIZ / "saida" / f"taxa_servico_{de or L['Data'].min()}_a_{ate or L['Data'].max()}.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        por_loja.to_excel(xw, sheet_name="Por loja", index=False)
        L.to_excel(xw, sheet_name="Lançamentos +10%", index=False)
        por_dia.to_excel(xw, sheet_name="Por dia", index=False)
        I.to_excel(xw, sheet_name="ISAIAS 01", index=False)
        if len(I_op):
            I_op.to_excel(xw, sheet_name="ISAIAS 01 por operador", index=False)
        for ws in xw.book.worksheets:
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
                ws.column_dimensions[get_column_letter(i)].width = min(42, max(10, max(len(str(c.value or "")) for c in col) + 2))
    print(por_loja.to_string(index=False))
    print(f"\nISAIAS 01: {len(I)} cupons · R$ {I['Serviço (R$)'].sum():.2f}")
    if len(I_op):
        print(I_op.to_string(index=False))
    print(f"gravado em {out}")


if __name__ == "__main__":
    main()
