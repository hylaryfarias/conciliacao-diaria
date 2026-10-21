"""Taxa de serviço não baixada, por loja, a partir do relatório de cupons (totais + operador).

    python3 scripts/isaias_taxa.py --loja "ISAIAS 01 - MARINGA" --de 2026-10-01 --ate 2026-10-20 \
        [--totais "entrada/pdv_totais_isa01_*.xlsx"] [--analitico "entrada/pdv_analitico_isa01_*.xlsx"] [--out saida/x.xlsx]

Não depende da forma de pagamento (serve para dias em que só há o relatório de cupons e o analítico):
1. casa cada venda aprovada da máquina Sicredi com o cupom de mesmo valor (janela: cupom de 30 min antes
   a 36 h depois da venda), o mais próximo no tempo;
2. nas vendas que sobram, procura cupom livre com total × 1,10 = valor da máquina (±1 centavo), na mesma
   janela: é o cupom fechado sem o item da taxa;
   (janela do +10%: cupom até 2 h depois da venda; o exato aceita até 36 h, por causa das comandas viradas);
3. para cada caso, anexa operador (cupons), atendente e mesa (analítico), se o cupom tem o item TAXA DE
   SERVICO, e a prova par a par (terminal exclusivo do caixa, minutos até o fechamento, concorrentes livres).
"""
import glob
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from concilia_pix import TAXA_SERVICO, carrega_lojas  # noqa: E402
from pdv_pix import HORA_VIRADA  # noqa: E402
from portal_sicredi import ler_portais  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
DIAS = {"Monday": "seg", "Tuesday": "ter", "Wednesday": "qua", "Thursday": "qui", "Friday": "sex", "Saturday": "sáb", "Sunday": "dom"}
FORMA = {"PIX": "PIX", "CRÉDITO": "Crédito", "CREDITO": "Crédito", "DÉBITO": "Débito", "DEBITO": "Débito"}


def ler_totais(padroes: str, loja: str) -> pd.DataFrame:
    fs = [f for p in padroes.split(",") for f in sorted(glob.glob(p))]
    partes = []
    for f in fs:
        d = pd.read_excel(f, engine="calamine")
        partes.append(d[d["Filial"].astype(str).str.strip() == loja])
    d = pd.concat(partes, ignore_index=True)
    if "Total Cupom" not in d.columns and "Total" in d.columns:   # analítico (itens): uma linha por cupom
        d["Caixa"] = d["Caixa"].astype(str).str.strip(); d["Cupom"] = d["Cupom"].astype(str).str.replace(".0", "", regex=False)
        d = d.groupby(["Filial", "Caixa", "Data", "Cupom"], as_index=False).agg(Hora=("Hora", "first"), **{"Total Cupom": ("Total", "first")}, Operador=("Operador", "first"))
    d["Caixa"] = d["Caixa"].astype(str).str.strip(); d["Cupom"] = d["Cupom"].astype(str).str.replace(".0", "", regex=False)
    d["Data"] = pd.to_datetime(d["Data"], dayfirst=True).dt.date
    d = d.drop_duplicates(["Caixa", "Data", "Cupom"])
    d["cent"] = (pd.to_numeric(d["Total Cupom"], errors="coerce").fillna(0) * 100).round().astype(int)
    d = d[d["cent"] > 0].copy()
    d["dt"] = pd.to_datetime(d["Data"].astype(str) + " " + d["Hora"].astype(str).str[:8])
    d.loc[d["dt"].dt.hour < HORA_VIRADA, "dt"] += pd.Timedelta(days=1)   # madrugada é o mesmo dia de movimento
    return d.sort_values("dt").reset_index(drop=True)


def ler_analitico(padroes: str, loja: str) -> pd.DataFrame:
    fs = [f for p in padroes.split(",") for f in sorted(glob.glob(p))]
    partes = []
    for f in fs:
        d = pd.read_excel(f, engine="calamine")
        partes.append(d[d["Filial"].astype(str).str.strip() == loja])
    if not partes:
        return pd.DataFrame()
    d = pd.concat(partes, ignore_index=True)
    d["Caixa"] = d["Caixa"].astype(str).str.strip(); d["Cupom"] = d["Cupom"].astype(str).str.replace(".0", "", regex=False)
    d["taxa"] = d["Produto"].astype(str).str.upper().str.contains("TAXA DE SERV")
    g = d.groupby(["Caixa", "Cupom"]).agg(Atendente=("Nome do atendente", "first"), Mesa=("Nr. mesa/comanda", "first"), taxa=("taxa", "any")).reset_index()
    g["Item taxa no cupom"] = g["taxa"].map({True: "sim", False: "não"})
    return g.drop(columns="taxa").drop_duplicates(["Caixa", "Cupom"])


MAQ_FORMAS = ("PIX MAQUININHA", "TEF - CREDITO", "TEF - DEBITO", "CARTAO CREDITO", "CARTAO DEBITO", "TEF - VOUCHER")


def ler_formas(padroes: str, loja: str) -> pd.DataFrame:
    """Relatório de cupons por forma de pagamento: só as linhas que passam na maquininha (Caixa, Cupom, Data, cent)."""
    fs = [f for p in padroes.split(",") for f in sorted(glob.glob(p))] if padroes else []
    partes = []
    for f in fs:
        d = pd.read_excel(f, engine="calamine")
        partes.append(d[d["Filial"].astype(str).str.strip() == loja])
    if not partes:
        return pd.DataFrame()
    d = pd.concat(partes, ignore_index=True)
    d["Caixa"] = d["Caixa"].astype(str).str.strip(); d["Cupom"] = d["Cupom"].astype(str).str.replace(".0", "", regex=False)
    d["Data"] = pd.to_datetime(d["Data"], dayfirst=True).dt.date
    d["cent"] = (pd.to_numeric(d["Vl. pagamento"], errors="coerce").fillna(0) * 100).round().astype(int)
    d["maq"] = d["Desc. pagam."].astype(str).str.upper().str.strip().isin(MAQ_FORMAS)
    return d[["Caixa", "Cupom", "Data", "cent", "maq"]]


def forma(m) -> str:
    p = str(m.get("Produto", "")).upper(); mod = str(m.get("Modalidade", "")).upper()
    if "PIX" in p or "PIX" in mod:
        return "PIX"
    for k, v in FORMA.items():
        if mod.startswith(k):
            return v
    return f"Voucher ({m.get('Bandeira', '')})"


def main():
    a = sys.argv[1:]
    g = lambda k, d=None: a[a.index(k) + 1] if k in a else d
    loja = g("--loja", "ISAIAS 01 - MARINGA")
    de = pd.Timestamp(g("--de", "2026-10-01")).date(); ate = pd.Timestamp(g("--ate")).date() if g("--ate") else None
    cod = "".join(w[0].lower() for w in loja.split()[:1]) + "".join(ch for ch in loja.split()[1] if ch.isdigit()) if len(loja.split()) > 1 else "loja"
    totais = g("--totais", str(RAIZ / "entrada" / f"pdv_totais_{'isa' if 'ISAIAS' in loja else ''}{cod[-2:]}_*.xlsx"))
    anal = g("--analitico", str(RAIZ / "entrada" / f"pdv_analitico_{'isa' if 'ISAIAS' in loja else ''}{cod[-2:]}_*.xlsx"))
    out = Path(g("--out", str(RAIZ / "saida" / f"isaias_taxa_servico_{de}_a_{ate or 'fim'}.xlsx")))
    portal = g("--portal", str(RAIZ / "entrada" / "portal" / "*.xlsx"))

    c = ler_totais(totais, loja)
    F = ler_formas(g("--formas", ""), loja)
    lojas = carrega_lojas()
    m = next(df for (f, _), df in ler_portais([f for p in portal.split(",") for f in glob.glob(p)], lojas).items() if f == loja)
    m = m[m["Status"].astype(str).str.upper().isin({"APROVADA", "AUTORIZADA"})].copy()
    m["cent"] = m["cent"].astype(int)
    m = m.sort_values("dt").reset_index(drop=True)
    if ate:
        c = c[c["Data"] <= ate]; m = m[m["Data"] <= ate]
    c = c[c["Data"] >= de]; m = m[m["Data"] >= de]
    print(f"{loja}: {len(c)} cupons com valor · {len(m)} vendas aprovadas na máquina · {de} a {ate}")

    # Pares candidatos (venda da máquina × cupom) na janela [-30 min, +36 h]: exatos e "+10%" (cupom × 1,10 = máquina ±1 c).
    # Atribuição gulosa pelo tempo: o par mais próximo primeiro; o exato leva vantagem de 5 min sobre o +10% de mesmo instante.
    c = c.copy()
    if len(F):
        # cupom presente no relatório por forma: as unidades são as linhas que passam na máquina (um cupom pode ter dois cartões);
        # cupom só com dinheiro, online ou prazo não entra. Cupom ausente do relatório por forma: usa o total.
        u = c.reset_index().merge(F, on=["Caixa", "Cupom", "Data"], how="left", suffixes=("", "_f"))
        tem = u["cent_f"].notna()
        u.loc[tem, "cent"] = u.loc[tem, "cent_f"]
        u = u[~tem | u["maq"].fillna(False).astype(bool)].drop(columns=["cent_f", "maq"])
        u["unid"] = range(len(u))
    else:
        u = c.reset_index(); u["unid"] = range(len(u))
    u["cent"] = u["cent"].astype(int); u["cent_u"] = u["cent"]; u["cent110"] = (u["cent"] * (1 + TAXA_SERVICO)).round().astype(int)
    ex = m.reset_index().merge(u, on="cent", suffixes=("_m", "_c")); ex["tipo"] = "exato"
    tx = pd.concat([m.reset_index().merge(u.assign(cent=lambda t: t["cent110"] + k), on="cent", suffixes=("_m", "_c")) for k in (-1, 0, 1)]); tx["tipo"] = "taxa"
    pares = pd.concat([ex, tx], ignore_index=True)
    pares["delta"] = (pares["dt_c"] - pares["dt_m"]).dt.total_seconds() / 60
    # exato: comanda pode fechar até 36 h depois (virada de dia); +10%: a taxa é cobrada na hora do pagamento, o cupom fecha logo em seguida (até 2 h)
    pares = pares[(pares["delta"] >= -30) & (pares["delta"] <= 36 * 60) & ((pares["tipo"] == "exato") | (pares["delta"] <= 2 * 60))].copy()
    pares["ordem"] = pares["delta"].abs() + (pares["tipo"] == "taxa") * 5
    pares = pares.sort_values("ordem")
    usados_u, usados_m, usados_c, par, casos, vcent = set(), set(), set(), {}, [], {}
    for _, r in pares.iterrows():
        i, j, k = r["index_m"], r["index_c"], r["unid"]
        if i in usados_m or k in usados_u:
            continue
        usados_m.add(i); usados_u.add(k); usados_c.add(j)
        if r["tipo"] == "exato":
            par[i] = j
        else:
            casos.append((i, j)); vcent[(i, j)] = int(r["cent_u"])
    c = c.assign(cent_total=c["cent"])
    print(f"casamentos exatos: {len(par)} · cupons fechados sem a taxa: {len(casos)}")

    # terminal × caixa nos casamentos exatos (prova de exclusividade)
    tc = pd.DataFrame([(m.loc[i, "Número do terminal"], c.loc[j, "Caixa"]) for i, j in par.items()], columns=["Terminal", "Caixa"])
    tc = tc.groupby(["Terminal", "Caixa"]).size().unstack(fill_value=0) if len(tc) else pd.DataFrame()
    A = ler_analitico(anal, loja)
    linhas = []
    for i, j in casos:
        v, x = m.loc[i], c.loc[j].copy(); x["cent"] = vcent.get((i, j), x["cent"])
        t = str(v["Número do terminal"]); usos = tc.loc[t] if t in tc.index else None
        if usos is not None and usos.sum() > 0:
            cx_term = f"{usos.idxmax()[-6:]} em {usos.max()} de {usos.sum()} casamentos exatos"
            exclusivo = "sim" if usos.max() >= 0.95 * usos.sum() and usos.idxmax() == x["Caixa"] else "não"   # 95%: um casamento cruzado por coincidência de valor não derruba a prova
        else:
            cx_term, exclusivo = "sem casamento exato no período", "?"
        livres = c[(~c.index.isin(usados_c)) & (c["cent"] == v["cent"]) & (c["dt"] >= v["dt"] - pd.Timedelta(minutes=30)) & (c["dt"] <= v["dt"] + pd.Timedelta(hours=3))]
        livres_m = m[(~m.index.isin(usados_m)) & (m["cent"] == x["cent"]) & (m["dt"] >= x["dt"] - pd.Timedelta(hours=3)) & (m["dt"] <= x["dt"] + pd.Timedelta(minutes=30))]
        linhas.append({
            "Loja": loja, "Forma": forma(v), "Data": f"{x['Data']:%d/%m/%Y}", "Dia": DIAS[pd.Timestamp(x["Data"]).day_name()],
            "Hora PDV": f"{x['dt']:%H:%M}", "Cupom": x["Cupom"], "Caixa": x["Caixa"], "Operador do caixa": x.get("Operador", ""),
            "Valor PDV": round(x["cent"] / 100, 2), "Valor máquina": round(v["cent"] / 100, 2), "Cobrado a mais na máquina (R$)": round((v["cent"] - x["cent"]) / 100, 2),
            "Hora máquina": f"{v['dt']:%H:%M}", "Terminal": t, "Autorização": str(v["Cód. de autorização"]),
            "Terminal usado por": cx_term, "Terminal exclusivo deste caixa": exclusivo,
            "Minutos máquina → fechamento": round((x["dt"] - v["dt"]).total_seconds() / 60),
            "Outro cupom livre = valor máquina (turno)": "; ".join(f"{r['Caixa'][-6:]} cupom {r['Cupom']} {r['dt']:%d/%m %H:%M}" for _, r in livres.iterrows()) or "nenhum",
            "Outra venda máq livre = valor cupom (turno)": "; ".join(f"{r['dt']:%d/%m %H:%M} {r['Modalidade']} {r['Número do terminal']}" for _, r in livres_m.iterrows()) or "nenhuma",
            "Prova": "fechada" if exclusivo == "sim" and not len(livres) and not len(livres_m) else "ver observação",
        })
    I = pd.DataFrame(linhas)
    if len(A) and len(I):
        I = I.merge(A, on=["Caixa", "Cupom"], how="left")
        cols = list(I.columns); cols.remove("Atendente"); cols.remove("Mesa"); cols.remove("Item taxa no cupom")
        k = cols.index("Operador do caixa") + 1
        I = I[cols[:k] + ["Atendente", "Mesa"] + cols[k:-0 or None] + ["Item taxa no cupom"]] if False else I[cols[:k] + ["Atendente", "Mesa"] + cols[k:] + ["Item taxa no cupom"]]
    I["_d"] = pd.to_datetime(I["Data"], dayfirst=True); I = I.sort_values(["_d", "Hora PDV"]).drop(columns="_d").reset_index(drop=True)
    por_op = I.groupby("Operador do caixa").agg(Cupons=("Cupom", "count"), **{"Sobra (R$)": ("Cobrado a mais na máquina (R$)", "sum")}).reset_index().sort_values("Sobra (R$)", ascending=False)
    por_dia = I.groupby(["Data", "Dia"], sort=False).agg(Cupons=("Cupom", "count"), **{"Sobra (R$)": ("Cobrado a mais na máquina (R$)", "sum")}).reset_index()
    noite = I[I["Hora PDV"] >= "20:00"]
    print(f"{len(I)} cupons · R$ {I['Cobrado a mais na máquina (R$)'].sum():.2f} · {len(noite)} entre 20h e 23h59 · prova: {I['Prova'].value_counts().to_dict()}")
    print(por_op.to_string(index=False)); print(por_dia.to_string(index=False))
    aba = " ".join(loja.split()[:2])
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        I.to_excel(xw, sheet_name=aba, index=False)
        por_op.to_excel(xw, sheet_name=f"{aba} por operador", index=False)
        por_dia.to_excel(xw, sheet_name=f"{aba} por dia", index=False)
        for ws in xw.book.worksheets:
            for k, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), 1):
                ws.column_dimensions[get_column_letter(k)].width = min(40, max(10, max(len(str(x.value or "")) for x in col) + 2))
            ws.freeze_panes = "A2"
    print("gravado em", out)


if __name__ == "__main__":
    main()
